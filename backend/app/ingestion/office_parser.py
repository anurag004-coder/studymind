"""Office and plain text document parsers (DOCX, PPTX, TXT/MD)."""

import io

import docx
from pptx import Presentation

from app.ingestion.models import ExtractedBlock, ParsedPage


class OfficeParser:
    """Parses Word (DOCX), PowerPoint (PPTX), and plain text files into ParsedPages."""

    @classmethod
    def parse_docx_bytes(cls, docx_bytes: bytes) -> list[ParsedPage]:
        """Parses a DOCX document into structured blocks with headings and tables."""
        doc = docx.Document(io.BytesIO(docx_bytes))
        blocks: list[ExtractedBlock] = []

        # Process paragraphs
        for p_idx, p in enumerate(doc.paragraphs):
            text = p.text.strip()
            if not text:
                continue

            style_name = p.style.name if p.style else ""
            block_type = "heading" if style_name.lower().startswith("heading") else "paragraph"

            blocks.append(
                ExtractedBlock(
                    page_no=1,
                    bbox=[50.0, float(50 + p_idx * 20), 550.0, float(70 + p_idx * 20)],
                    block_type=block_type,  # type: ignore[arg-type]
                    text=text,
                    reading_order=len(blocks),
                )
            )

        # Process tables
        for table in doc.tables:
            rows_data: list[list[str]] = []
            for row in table.rows:
                row_cells = [cell.text.strip() for cell in row.cells]
                rows_data.append(row_cells)

            if rows_data:
                # Format to Markdown table
                headers = rows_data[0]
                divider = ["---"] * len(headers)
                md_lines = [" | ".join(headers), " | ".join(divider)]
                for r in rows_data[1:]:
                    md_lines.append(" | ".join(r))

                table_md = "\n".join(md_lines)
                blocks.append(
                    ExtractedBlock(
                        page_no=1,
                        bbox=[50.0, 400.0, 550.0, 550.0],
                        block_type="table",
                        text=table_md,
                        reading_order=len(blocks),
                    )
                )

        if not blocks:
            blocks.append(
                ExtractedBlock(
                    page_no=1,
                    bbox=[50.0, 50.0, 550.0, 100.0],
                    block_type="paragraph",
                    text="[Empty Document]",
                    reading_order=0,
                )
            )

        return [
            ParsedPage(
                page_no=1,
                has_text_layer=True,
                blocks=blocks,
                raw_text="\n\n".join(b.text for b in blocks),
                ocr_confidence=1.0,
            )
        ]

    @classmethod
    def parse_pptx_bytes(cls, pptx_bytes: bytes) -> list[ParsedPage]:
        """Parses a PPTX presentation where each slide is treated as a page."""
        prs = Presentation(io.BytesIO(pptx_bytes))
        pages: list[ParsedPage] = []

        for slide_idx, slide in enumerate(prs.slides):
            page_no = slide_idx + 1
            blocks: list[ExtractedBlock] = []

            for shape_idx, shape in enumerate(slide.shapes):
                # Text frames
                if shape.has_text_frame:
                    text = shape.text_frame.text.strip()
                    if text:
                        is_title = shape == slide.shapes.title
                        blocks.append(
                            ExtractedBlock(
                                page_no=page_no,
                                bbox=[
                                    float(shape.left or 50),
                                    float(shape.top or 50),
                                    float((shape.left or 50) + (shape.width or 400)),
                                    float((shape.top or 50) + (shape.height or 100)),
                                ],
                                block_type="heading" if is_title else "paragraph",
                                text=text,
                                reading_order=shape_idx,
                            )
                        )

                # Tables
                elif shape.has_table:
                    table = shape.table
                    rows: list[list[str]] = []
                    for row in table.rows:
                        rows.append([cell.text.strip() for cell in row.cells])
                    if rows:
                        headers = rows[0]
                        md_table = " | ".join(headers) + "\n" + " | ".join(["---"] * len(headers))
                        for r in rows[1:]:
                            md_table += "\n" + " | ".join(r)
                        blocks.append(
                            ExtractedBlock(
                                page_no=page_no,
                                bbox=[50.0, 200.0, 500.0, 400.0],
                                block_type="table",
                                text=md_table,
                                reading_order=shape_idx,
                            )
                        )

            pages.append(
                ParsedPage(
                    page_no=page_no,
                    has_text_layer=True,
                    blocks=blocks,
                    raw_text="\n\n".join(b.text for b in blocks),
                    ocr_confidence=1.0,
                )
            )

        return pages

    @classmethod
    def parse_text_bytes(cls, text_bytes: bytes, filename: str = "notes.txt") -> list[ParsedPage]:
        """Parses plain text / Markdown bytes into structured pages."""
        text = text_bytes.decode("utf-8", errors="replace").strip()
        lines = text.splitlines()

        blocks: list[ExtractedBlock] = []
        current_para: list[str] = []

        def flush_para() -> None:
            if current_para:
                para_text = " ".join(current_para).strip()
                if para_text:
                    blocks.append(
                        ExtractedBlock(
                            page_no=1,
                            bbox=[
                                50.0,
                                float(50 + len(blocks) * 30),
                                550.0,
                                float(80 + len(blocks) * 30),
                            ],
                            block_type="paragraph",
                            text=para_text,
                            reading_order=len(blocks),
                        )
                    )
                current_para.clear()

        for line in lines:
            stripped = line.strip()
            if not stripped:
                flush_para()
                continue

            if stripped.startswith(("#", "Chapter", "Section")):
                flush_para()
                blocks.append(
                    ExtractedBlock(
                        page_no=1,
                        bbox=[
                            50.0,
                            float(50 + len(blocks) * 30),
                            550.0,
                            float(80 + len(blocks) * 30),
                        ],
                        block_type="heading",
                        text=stripped.lstrip("#").strip(),
                        reading_order=len(blocks),
                    )
                )
            else:
                current_para.append(stripped)

        flush_para()

        return [
            ParsedPage(
                page_no=1,
                has_text_layer=True,
                blocks=blocks,
                raw_text=text,
                ocr_confidence=1.0,
            )
        ]
